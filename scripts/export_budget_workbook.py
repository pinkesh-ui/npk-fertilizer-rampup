"""Export equal-pace budgets / CAPEX to a multi-tab Excel workbook."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from agricultural_input_rampup import (  # noqa: E402
    CAPEX_FUNCS,
    COMMODITIES,
    COST_CAPACITY_EXPONENT,
    DEFAULT_FRACTION_FUNCTIONING,
    DEFAULT_STARTUP_PCT_OF_FULL,
    FAST_CONSTRUCTION_COST_FACTOR,
    FERTILIZER_KEYS,
    NH3_TARGET_PLANT_SIZE_TPD,
    PESTICIDE_KEYS,
    REFERENCE_ANNUAL_BUDGET_USD,
    allocate_equal_pace_budgets,
    category_budget_totals,
)

HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="1F4E79")
SECTION_FONT = Font(name="Calibri", bold=True, size=12, color="1F4E79")
NORMAL = Font(name="Calibri", size=11)
BOLD = Font(name="Calibri", bold=True, size=11)
HEADER_FILL = PatternFill("solid", fgColor="2F5D50")
SECTION_FILL = PatternFill("solid", fgColor="D6EAF8")
ALT_FILL = PatternFill("solid", fgColor="F4F6F7")
THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)

ORDER = [
    "nh3",
    "phosphate",
    "potassium",
    "herbicide",
    "fungicide",
    "insecticide",
    "h2so4",
]
PEST_TPY = {"herbicide": 600.0, "insecticide": 2100.0, "fungicide": 1100.0}


def style_header(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN


def autosize(ws, min_w: int = 12, max_w: int = 46) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col:
            if cell.value is None:
                continue
            length = max(length, len(str(cell.value)))
        ws.column_dimensions[letter].width = max(min_w, min(max_w, length + 3))


def main() -> Path:
    budgets = allocate_equal_pace_budgets()
    cats = category_budget_totals()
    total = REFERENCE_ANNUAL_BUDGET_USD
    by_key = {c.key: c for c in COMMODITIES}
    cat_of = {k: "Fertilizer" for k in FERTILIZER_KEYS}
    cat_of.update({k: "Pesticides" for k in PESTICIDE_KEYS})
    cat_of["h2so4"] = "H2SO4"
    cat_totals = {
        "Fertilizer": cats["fertilizer"],
        "Pesticides": cats["pesticides"],
        "H2SO4": cats["h2so4"],
    }

    wb = Workbook()

    # --- Summary ---
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Agricultural Input Ramp-up — Equal-pace Budget Summary"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")
    ws["A3"] = "Shared construction pot (USD/yr)"
    ws["B3"] = total
    ws["B3"].number_format = '"$"#,##0'
    ws["A4"] = "Allocation rule"
    ws["B4"] = "B_i proportional to (CAPEX $/tpa) x (current t/yr)"
    ws["A5"] = "Startup % of Fully Scaled Production"
    ws["B5"] = DEFAULT_STARTUP_PCT_OF_FULL
    ws["A6"] = "fraction_functioning_after_disruption"
    ws["B6"] = DEFAULT_FRACTION_FUNCTIONING

    ws["A8"] = "Category totals"
    ws["A8"].font = SECTION_FONT
    for i, h in enumerate(
        ["Category", "Budget (USD/yr)", "Budget ($B/yr)", "Share of pot (%)"], 1
    ):
        ws.cell(row=9, column=i, value=h)
    style_header(ws, 9, 4)

    for r, (name, val) in enumerate(
        [
            ("Fertilizer", cats["fertilizer"]),
            ("Pesticides", cats["pesticides"]),
            ("H2SO4", cats["h2so4"]),
            ("TOTAL", total),
        ],
        10,
    ):
        ws.cell(row=r, column=1, value=name).font = BOLD if name == "TOTAL" else NORMAL
        ws.cell(row=r, column=2, value=val).number_format = '"$"#,##0'
        ws.cell(row=r, column=3, value=val / 1e9).number_format = "0.00"
        ws.cell(row=r, column=4, value=100 * val / total).number_format = "0.00"
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = THIN
            if r % 2 == 0:
                ws.cell(row=r, column=c).fill = ALT_FILL

    ws["A15"] = "Notes"
    ws["A15"].font = SECTION_FONT
    ws["A16"] = (
        "NH3 is the reference template for ramp-up timing/charts; "
        "the $758B pot is shared across all commodities."
    )
    ws["A17"] = (
        "Pesticide current use is 2026 Mt active ingredient / year; "
        "pesticide plant sizes are in t/yr."
    )
    ws["A18"] = "Generated from the npk-fertilizer-rampup equal-pace allocator."
    autosize(ws)

    # --- Budgets ---
    ws = wb.create_sheet("Budgets by commodity")
    ws["A1"] = "Equal-pace annual construction budgets"
    ws["A1"].font = TITLE_FONT
    headers = [
        "Category",
        "Commodity",
        "Key",
        "Current (Mt/yr)",
        "Budget (USD/yr)",
        "Budget ($B/yr)",
        "Share of pot (%)",
        "Share of category (%)",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)
    style_header(ws, 3, len(headers))
    ws.row_dimensions[3].height = 30

    for r, key in enumerate(ORDER, 4):
        com = by_key[key]
        cat = cat_of[key]
        bud = budgets[key]
        vals = [
            cat,
            com.label,
            key,
            com.current_mt_per_year,
            bud,
            bud / 1e9,
            100 * bud / total,
            100 * bud / cat_totals[cat],
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = THIN
            cell.font = NORMAL
            if c == 5:
                cell.number_format = '"$"#,##0'
            elif c in (4, 6, 7, 8):
                cell.number_format = "0.00"
            if r % 2 == 0:
                cell.fill = ALT_FILL
    autosize(ws)

    # --- CAPEX ---
    ws = wb.create_sheet("CAPEX & plant size")
    ws["A1"] = "CAPEX benchmarks and plant sizes"
    ws["A1"].font = TITLE_FONT
    headers = [
        "Category",
        "Commodity",
        "Key",
        "Target plant size",
        "Plant size unit",
        "Regular CAPEX (USD/plant)",
        "Fast CAPEX (USD/plant)",
        "Capital efficiency ($/tpa)",
        "Weight = $/tpa x t/yr",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)
    style_header(ws, 3, len(headers))
    ws.row_dimensions[3].height = 32

    for r, key in enumerate(ORDER, 4):
        com = by_key[key]
        capex = CAPEX_FUNCS[key]()
        current_tpy = com.current_mt_per_year * 1_000_000.0
        weight = capex.capital_efficiency_usd_per_tpa * current_tpy
        if key in PEST_TPY:
            size, unit = PEST_TPY[key], "t/yr"
        elif key == "nh3":
            size, unit = NH3_TARGET_PLANT_SIZE_TPD, "t/day"
        else:
            size, unit = com.target_plant_size_tpd, "t/day"
        vals = [
            cat_of[key],
            com.label,
            key,
            size,
            unit,
            capex.regular_capex_usd,
            capex.fast_capex_usd,
            capex.capital_efficiency_usd_per_tpa,
            weight,
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = THIN
            if c in (6, 7, 9):
                cell.number_format = '"$"#,##0.00'
            elif c in (4, 8):
                cell.number_format = "0.00"
            if r % 2 == 0:
                cell.fill = ALT_FILL
    autosize(ws)

    # --- By category ---
    ws = wb.create_sheet("By category")
    ws["A1"] = "Budgets grouped by category"
    ws["A1"].font = TITLE_FONT
    row = 3
    for cat_name, keys in [
        ("Fertilizer", FERTILIZER_KEYS),
        ("Pesticides", PESTICIDE_KEYS),
        ("H2SO4", ("h2so4",)),
    ]:
        ws.cell(row=row, column=1, value=cat_name).font = SECTION_FONT
        row += 1
        for i, h in enumerate(
            ["Commodity", "Budget ($B/yr)", "Share of category (%)", "Share of pot (%)"],
            1,
        ):
            ws.cell(row=row, column=i, value=h)
        style_header(ws, row, 4)
        row += 1
        cat_total = sum(budgets[k] for k in keys)
        for key in keys:
            com = by_key[key]
            bud = budgets[key]
            ws.cell(row=row, column=1, value=com.label).border = THIN
            ws.cell(row=row, column=2, value=bud / 1e9).number_format = "0.00"
            ws.cell(row=row, column=2).border = THIN
            ws.cell(row=row, column=3, value=100 * bud / cat_total).number_format = "0.00"
            ws.cell(row=row, column=3).border = THIN
            ws.cell(row=row, column=4, value=100 * bud / total).number_format = "0.00"
            ws.cell(row=row, column=4).border = THIN
            row += 1
        ws.cell(row=row, column=1, value=f"{cat_name} total").font = BOLD
        ws.cell(row=row, column=2, value=cat_total / 1e9).number_format = "0.00"
        ws.cell(row=row, column=2).font = BOLD
        for c in range(1, 5):
            ws.cell(row=row, column=c).border = THIN
            ws.cell(row=row, column=c).fill = SECTION_FILL
        row += 2
    autosize(ws)

    # --- Assumptions ---
    ws = wb.create_sheet("Assumptions")
    ws["A1"] = "Shared model assumptions"
    ws["A1"].font = TITLE_FONT
    for i, h in enumerate(["Parameter", "Value", "Notes"], 1):
        ws.cell(row=3, column=i, value=h)
    style_header(ws, 3, 3)

    assumptions = [
        ("REFERENCE_ANNUAL_BUDGET_USD", total, "Shared world construction pot ($/yr)"),
        (
            "DEFAULT_STARTUP_PCT_OF_FULL",
            DEFAULT_STARTUP_PCT_OF_FULL,
            "New plant starts at this share of full output",
        ),
        (
            "DEFAULT_FRACTION_FUNCTIONING",
            DEFAULT_FRACTION_FUNCTIONING,
            "Surviving existing + new-build rate after disruption",
        ),
        ("COST_CAPACITY_EXPONENT", COST_CAPACITY_EXPONENT, "Plant cost scaling exponent"),
        (
            "FAST_CONSTRUCTION_COST_FACTOR",
            FAST_CONSTRUCTION_COST_FACTOR,
            "Fast CAPEX = factor x regular CAPEX",
        ),
        (
            "NH3_TARGET_PLANT_SIZE_TPD",
            NH3_TARGET_PLANT_SIZE_TPD,
            "EuroChem Kingisepp-scaled NH3 plant (t/day)",
        ),
        (
            "NH3 regular CAPEX",
            1_368_630_195.13,
            "CAPEX Calcs B19 — Using this estimate (USD)",
        ),
        (
            "NH3 USD inflation conversion (2025)",
            1.33,
            "Sheet: USD inflation conversion",
        ),
        ("Herbicide current use 2026", 1.9902, "Mt ai/yr"),
        ("Insecticide current use 2026", 0.8232, "Mt ai/yr"),
        ("Fungicide current use 2026", 0.8624, "Mt ai/yr"),
        ("Herbicide target plant", 600, "t/yr; CAPEX $16,593,135"),
        ("Insecticide target plant", 2100, "t/yr; CAPEX $12,900,141"),
        ("Fungicide target plant", 1100, "t/yr; CAPEX $31,323,634"),
    ]
    for r, (param, value, notes) in enumerate(assumptions, 4):
        ws.cell(row=r, column=1, value=param).border = THIN
        cell = ws.cell(row=r, column=2, value=value)
        cell.border = THIN
        if isinstance(value, (int, float)) and value >= 1000:
            cell.number_format = "#,##0.##"
        ws.cell(row=r, column=3, value=notes).border = THIN
        if r % 2 == 0:
            for c in range(1, 4):
                ws.cell(row=r, column=c).fill = ALT_FILL
    autosize(ws)

    out_dir = ROOT / "results"
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / "agricultural_input_equal_pace_budgets.xlsx"
    wb.save(out_path)

    downloads = Path.home() / "Downloads" / "agricultural_input_equal_pace_budgets.xlsx"
    wb.save(downloads)
    print(out_path)
    print(downloads)
    return out_path


if __name__ == "__main__":
    main()
